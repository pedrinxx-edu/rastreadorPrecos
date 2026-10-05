import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from docx import Document

def procurar_preco():
    repeticao = "sim"
    while repeticao == "sim":
        precoMin = float(input("Qual o preço mínimo para filtrar? "))
        precoMax = float(input("Qual o preço máximo para filtrar? "))

        if precoMin < precoMax:
            with open ('catalogoCompleto.csv', mode='r', encoding='utf-8') as arquivo:
                leitor = csv.DictReader(arquivo)
                for linha in leitor:
                    preco_sem_simbolo = linha['valor'].replace('£', '')
                    precoLimpo = float(preco_sem_simbolo)

                    if precoMin <= precoLimpo <= precoMax:
                        print(f"| {linha['nome_do_livro']} | {linha['valor']} |")
            repeticao = "não"
        else:
            print("Erro: O valor máximo tem que ser maior que o valor mínimo!")
            repeticao = "sim"

def procurar_titulo():
    repeticao = "sim"

    while repeticao == "sim":
        palavraChave = input("Digite a palavra-chave para procurar no título do livro: ").strip() .lower()
        encontrou = False

        with open ('catalogoCompleto.csv', mode='r', encoding='utf-8') as arquivo:
            leitor = csv.DictReader(arquivo)
            for linha in leitor:
                tituloFormatado = linha['nome_do_livro'].strip() .lower()

                if palavraChave in tituloFormatado:
                    print(f"| {linha['nome_do_livro']} | {linha['valor']} |")
                    encontrou = True
        if encontrou == False:
            print("A palavra-chave não tem semelhança com nenhum título, tente novamente!")
        else:
            repeticao = "não"

def exportar_catalogo_excel():
    print("\nGerando seu arquivo Excel, aguarde...")

    planilha = Workbook()
    aba = planilha.active
    aba.title = "Catálogo de Livros"

    aba['A1'] = "Nome do Livro"
    aba['B1'] = "Preço (£)"

    estilo_fonte = Font(bold=True, color="FFFFFF")
    estilo_fundo = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")

    aba['A1'].font = estilo_fonte
    aba['A1'].fill = estilo_fundo
    aba['B1'].font = estilo_fonte
    aba['B1'].fill = estilo_fundo

    aba.column_dimensions['A'].width = 60
    aba.column_dimensions['B'].width = 15

    with open ('catalogoCompleto.csv', mode='r', encoding='utf-8') as arquivo:
        leitor = csv.DictReader(arquivo)

        linha_excel = 2

        for linha in leitor:
            aba[f'A{linha_excel}'] = linha['nome_do_livro'].strip()

            preco_limpo = float(linha['valor'].replace('£', ''))
            aba[f'B{linha_excel}'] = preco_limpo

            linha_excel += 1

    planilha.save("catalogo_formatado.xlsx")
    print("Arquivo Excel gerado com sucesso! O arquivo foi salvo como 'catalogo_formatado.xlsx'")

def exportar_catalogo_word():
    print("\nGerando seu arquivo Word, aguarde...")

    documento = Document()

    documento.add_heading("Catálogo de livros - Books to Scraoe", level=1)
    documento.add_paragraph("Abaixo está a lista completa de livros extraídos: \n")

    with open('catalogoCompleto.csv', mode='r', encoding='utf-8') as arquivo:
        leitor = csv.DictReader(arquivo)

        for linha in leitor:
            nome_limpo = linha['nome_do_livro'].strip()
            preco = linha['valor']

            paragrafo = documento.add_paragraph()

            texto_nome = paragrafo.add_run(f"Livro: {nome_limpo}")
            texto_nome.bold = True

            paragrafo.add_run(f"\nPreço: {preco}\n")
    documento.save("Catalogo_de_livros.docx")
    print("\nArquivo Word gerado com sucesso! O arquivo foi salvo como 'Catalogo_de_livros.docx'")